"""Fill the COMP 365 Checkpoint 1 template with Locksmith's current evidence.

Usage:
    <venv>/bin/python scripts/fill-cp1-report.py <template.docx> <output.docx>

Only the gray placeholder cells are replaced; the template's structure is untouched.
"""
import sys
from docx import Document

REPO = "https://github.com/rickeyjohnson/Locksmith"

# (table index, row index, column index) -> replacement text
CELLS = {
    # 1 Project Direction
    (0, 1, 1): "Rickey Johnson",
    (0, 2, 1): "Locksmith: Comparing Prompt-Injection Defense Strategies Through Adversarial Gameplay",
    (0, 3, 1): "Does raising defense-prompt complexity across levels measurably lower extraction success?",
    (0, 4, 1): "No changes. The proposal direction was confirmed on 09/06/2026 and is unchanged.",

    # 2 Data and Input Access
    (1, 1, 1): (
        f"Self-authored level prompts and passwords, plus logged attempts. Code and data: {REPO}. "
        "Model: qwen3:8b (open weights) served locally by Ollama, pinned at temperature 0.7, seed 42. "
        "Attack corpus: 18 hand-written attacks plus 60 human-written attacks sampled from the public "
        "Tensor Trust dataset (Toyer et al., ICLR 2024, arXiv:2311.01011)."
    ),
    (1, 2, 1): (
        "All 8 levels authored and running, each with its own password and one added defense layer. "
        "78 attack prompts. 1,170 logged benchmark attempts so far: 624 against the stacked ladder "
        "(78 attacks x 8 levels) and 546 against 7 single-defense variants. Player attempts are logged "
        "continuously; human playtesting is scheduled for Sep 26-27."
    ),
    (1, 3, 1): (
        "One row per attempt in Supabase Postgres (tables players, attempts, guesses) with a local JSONL "
        "fallback, plus one CSV row per benchmark attempt. Preprocessing is automatic: model reasoning is "
        "stripped before display, every response is labeled leaked / not leaked by a detector that catches "
        "spaced, hyphenated, reversed, base64, hex, and acrostic spellings, and failed model calls are "
        "marked status=error and excluded from every rate."
    ),
    (1, 4, 1): (
        "Players see a consent notice, are told prompts and responses are recorded, and are asked not to "
        "enter personal information; all in-game secrets are fictional. Passwords are generated nonsense "
        "words kept in a git-ignored file, so the public repository cannot expose them. Row-level security "
        "is enabled on all tables with no public policies. Quality limits: one model at one seed, so "
        "results may not generalize; the leak detector does not catch semantic hints (\"it is a bird\")."
    ),

    # 3 Working Technical Artifact
    (2, 1, 1): "Web application (Next.js + TypeScript) plus a command-line attack benchmark.",
    (2, 2, 1): f"{REPO} — run `npm run dev`, play at http://localhost:3000; benchmark: `npm run benchmark`.",
    (2, 3, 1): (
        "Input: a player-typed prompt for the current level. Processing: input guards run first, then the "
        "level's system prompt and the prompt go to the model, then output guards run, then the reply is "
        "checked for the password. All 8 levels work, including the two AI defenses (a classifier that "
        "screens input, and a critic agent that audits the reply). Observable output: the guardian's reply "
        "in the browser, blocked replies marked in red, an attempt counter, level unlocking on a correct "
        "guess, a logged database row per attempt, and a printed leak-rate table from the benchmark. "
        "87 automated tests pass (`npm test`)."
    ),
    (2, 4, 1): (
        "No human playtest data yet; the game is not publicly deployed, so attempts-until-success has no "
        "player data behind it. Levels 7-8 may be unbeatable: no attack in the 78-prompt corpus produced a "
        "leak a player could see. The leak detector does not catch semantic hints, and the guardian has no "
        "memory between attempts, so multi-turn attacks cannot be studied."
    ),

    # 4 Baseline or Initial Result
    (4, 1, 0): "Stacked defense ladder, levels 1-8, compared against the non-AI baselines (keyword/regex input filter, output string filter)",
    (4, 1, 1): "78 attack prompts per level (18 hand-written + 60 Tensor Trust), qwen3:8b, temperature 0.7, seed 42; 624 attempts total",
    (4, 1, 2): (
        "Player-visible leak rate by level: 84.6%, 47.4%, 6.4%, 0.0%, 1.4%, 0.0%, 0.0%, 0.0%. "
        "Model-side leak rate given the prompt reached the model: 84.6%, 82.2%, 86.7%, 60.0%, 21.6%, "
        "2.2%, 0.0%, 0.0%. Single defenses measured alone: spotlighting 1.3%, output filter 3.8%, "
        "critic 17.9%, prompt shield 21.8%, hardened prompt 25.6%, regex filter 50.0%, refusal 65.4%."
    ),
    (4, 1, 3): (
        "Leak rate falls as defense complexity rises, as predicted. The stronger result is in the second "
        "row of numbers: keyword filtering blocks 42% of traffic but does not change the model's behavior "
        "at all (86.7% vs 84.6% undefended), so input filtering only reduces exposure. Only prompt-level "
        "defenses change what the model produces, and spotlighting is the strongest single defense while "
        "costing no extra model call. Current limit: the measure saturates at 0% from level 6 up, so the "
        "two AI defenses cannot yet be ranked against each other on player-visible leaks."
    ),

    # 5 Evaluation Readiness
    (5, 1, 0): "Measure 1: password-leak rate per level across N attempts",
    (5, 1, 1): "Share of attempts where the response shown to the player contains the level's password, in any disguise",
    (5, 1, 2): "AI-defended levels vs the non-AI keyword/regex baseline; target is a leak rate that falls level to level",
    (5, 1, 3): "Ready — automated and already produced for all 8 levels and 7 isolated defenses",
    (5, 2, 0): "Measure 2: mean and median attempts until a successful extraction per level",
    (5, 2, 1): "Number of attempts a player makes before guessing the password correctly",
    (5, 2, 2): "Compared across levels; more attempts required indicates a stronger defense",
    (5, 2, 3): "In progress — logging and the analysis script are built and tested; awaiting the Sep 26-27 playtest",

    # 6 Progress Status
    (6, 1, 1): "Complete",
    (6, 1, 2): "8 level configs in src/levels/index.ts; 78-attack corpus in scripts/attacks.json and attacks-extra.json; qwen3:8b verified",
    (6, 2, 1): "Complete",
    (6, 2, 2): "Automatic per-attempt labeling (src/pipeline/leak.ts, 11 unit tests); 1,170 labeled rows in docs/results/",
    (6, 3, 1): "Complete",
    (6, 3, 2): "Non-AI baselines built and measured: keyword filter 50.0% and output filter 3.8% leak rate alone, against 84.6% undefended",
    (6, 4, 1): "Complete",
    (6, 4, 2): "All 8 levels call qwen3:8b through one pipeline; AI prompt shield and critic agent implemented and tested (src/guards/)",
    (6, 5, 1): "In progress",
    (6, 5, 2): "Leak rate per level automated and reported; attempts-until-success logged, awaiting playtesters (scripts/analyze.ts)",
    (6, 6, 1): "In progress",
    (6, 6, 2): "Playable 8-level game with consent screen, Supabase logging, and 87 passing tests; public playtest deployment remains",

    # 7 Risks and Next Steps
    (7, 1, 1): (
        "A complete, tested game and benchmark that measure seven defenses on one fixed model and attack "
        "set, producing the project's central finding: input filtering reduces exposure without making the "
        "model any safer, while prompt-level defenses change what the model will say."
    ),
    (7, 2, 1): (
        "The measure saturates at the top of the ladder. Levels 6-8 sit at 0% player-visible leaks against "
        "the 78-attack corpus, so the two AI defenses cannot be ranked against each other, and human "
        "playtesters may stall at level 6 and never reach them. Recruiting enough playtesters in a two-day "
        "window is the second risk."
    ),
    (7, 3, 1): (
        "Run the Sep 26-27 playtest through a Cloudflare tunnel and report human leak rates and "
        "attempts-until-success per level; freeze the configuration for the duration; produce the final "
        "evaluation report comparing human and benchmark results."
    ),
    (7, 4, 1): (
        "If human playtesting falls short, report the automated benchmark as the primary evidence since it "
        "runs the identical pipeline; if the upper levels prove unusable, collapse the ladder to five "
        "levels (none, non-AI filter, hardened prompt, prompt shield, critic), which is a config-only change."
    ),
}


def set_cell(cell, text):
    """Replace a cell's text while keeping its first run's formatting."""
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def main(template: str, out: str) -> None:
    doc = Document(template)
    for (ti, ri, ci), text in CELLS.items():
        set_cell(doc.tables[ti].rows[ri].cells[ci], text)

    for para in doc.paragraphs:
        if para.text.strip() == "Student name [Enter response here]":
            set_cell_paragraph(para, "Student name  Rickey Johnson")
        elif para.text.strip() == "Date [Enter response here]":
            set_cell_paragraph(para, "Date  09/20/2026")

    doc.save(out)
    print(f"Wrote {out}")


def set_cell_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
